import sys
from pathlib import Path
import fitz
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pdf2docx import Converter
def main():
    if len(sys.argv) != 3:
        raise RuntimeError("Expected input PDF and output DOCX paths")
    source = Path(sys.argv[1]).resolve()
    target = Path(sys.argv[2]).resolve()
    if not source.is_file() or source.suffix.lower() != ".pdf":
        raise RuntimeError("Invalid PDF input")
    pdf = fitz.open(str(source))
    try:
        text = "\n".join(page.get_text() for page in pdf)
        label_mode = (
            "ARSH GLASS LED Mirror" in text
            or text.count("X002GKH599") >= 3
        )
        if label_mode:
            doc = Document()
            section = doc.sections[0]
            first_page = pdf[0].rect
            section.page_width = Inches(first_page.width / 72)
            section.page_height = Inches(first_page.height / 72)
            section.top_margin = Inches(0)
            section.bottom_margin = Inches(0)
            section.left_margin = Inches(0)
            section.right_margin = Inches(0)
            section.header_distance = Inches(0)
            section.footer_distance = Inches(0)
            for index, page in enumerate(pdf):
                if index:
                    doc.add_page_break()
                paragraph = doc.paragraphs[-1] if index == 0 else doc.paragraphs[-1]
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.left_indent = Inches(0)
                paragraph.paragraph_format.right_indent = Inches(0)
                image_path = target.parent / f"pdfword_page_{index + 1}.png"
                pix = page.get_pixmap(matrix=fitz.Matrix(3, 3), alpha=False)
                pix.save(str(image_path))
                run = paragraph.add_run()
                run.add_picture(
                    str(image_path),
                    width=Inches(page.rect.width / 72),
                    height=Inches(page.rect.height / 72),
                )
                image_path.unlink(missing_ok=True)
            doc.save(str(target))
        else:
            converter = Converter(str(source))
            try:
                converter.convert(str(target))
            finally:
                converter.close()
    finally:
        pdf.close()
    if not target.is_file() or target.stat().st_size == 0:
        raise RuntimeError("Conversion did not produce a DOCX file")
if __name__ == "__main__":
    main()
